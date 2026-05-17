from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Пояснительная_записка_СУБД_B_plus_tree.docx"


def set_font(run, name="Times New Roman", size=14, bold=False, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, text, fld_end])


def configure_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.different_first_page_header_footer = True
    add_page_number(section.footer.paragraphs[0])

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(14)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.first_line_indent = Cm(1.25)

    for name in ["Heading 1", "Heading 2", "Heading 3"]:
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(16)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.first_line_indent = Cm(0)

    return doc


def p(doc, text="", *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, bold=False, italic=False):
    par = doc.add_paragraph()
    par.alignment = align
    par.paragraph_format.line_spacing = 1.15
    par.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    par.paragraph_format.space_after = Pt(6)
    par.paragraph_format.first_line_indent = Cm(1.25) if indent else Cm(0)
    run = par.add_run(text)
    set_font(run, bold=bold, italic=italic)
    return par


def heading(doc, text, level=1):
    par = doc.add_paragraph(style=f"Heading {level}")
    run = par.add_run(text)
    set_font(run, size=16, bold=True)
    return par


def listing_caption(doc, number, text):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    par.paragraph_format.line_spacing = 1.0
    par.paragraph_format.space_before = Pt(8)
    par.paragraph_format.space_after = Pt(3)
    run = par.add_run(f"Листинг {number}. {text}")
    set_font(run, size=12, italic=True)


def code(doc, text):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    par.paragraph_format.line_spacing = 1.0
    par.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    par.paragraph_format.space_after = Pt(6)
    par.paragraph_format.first_line_indent = Cm(0)
    for line in text.strip("\n").splitlines():
        run = par.add_run(line + "\n")
        set_font(run, name="Consolas", size=12)


def add_title_page(doc):
    for line in [
        "Московский авиационный институт",
        "(Национальный исследовательский университет)",
        'Институт № 8 «Компьютерные науки и прикладная математика»',
        'Кафедра 806 «Вычислительная математика и программирование»',
    ]:
        p(doc, line, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)

    for _ in range(5):
        doc.add_paragraph()

    p(doc, "Курсовая работа", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, bold=True)
    p(doc, "по курсу", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    p(doc, "«Системное программирование»", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    p(doc, "4 семестр", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    p(
        doc,
        "«Разработка СУБД с SQL-подобным языком запросов и индексом на основе B+-дерева»",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        indent=False,
        bold=True,
    )

    for _ in range(5):
        doc.add_paragraph()

    block = doc.add_paragraph()
    block.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    block.paragraph_format.line_spacing = 1.15
    block.paragraph_format.first_line_indent = Cm(0)
    run = block.add_run(
        "Выполнил: ______________________\n"
        "Группа: ________________________\n"
        "Преподаватель: _________________\n"
        "Оценка: ______\n"
        "Дата: ______"
    )
    set_font(run, size=14)

    for _ in range(4):
        doc.add_paragraph()
    p(doc, "Москва, 2026", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    doc.add_page_break()


def add_contents(doc):
    heading(doc, "Содержание", 1)
    items = [
        ("Введение", "3"),
        ("Основная часть", "4"),
        ("1. Базовая СУБД с индексом B+-tree (задание 0)", "4"),
        ("2. Поддержка темпоральной персистентности (задание 1)", "9"),
        ("3. Оптимизация хранения строковых данных (задание 2)", "11"),
        ("4. Логирование активности (задание 7)", "12"),
        ("5. Значения по умолчанию (задание 10)", "13"),
        ("6. Составные логические выражения (задание 11)", "14"),
        ("7. Агрегатные функции SUM, COUNT, AVG (задание 12)", "16"),
        ("Тестирование", "17"),
        ("Вывод", "19"),
        ("Список использованных источников", "20"),
        ("Приложение А. Пример сценария работы", "21"),
        ("Приложение Б. Ключевые фрагменты реализации", "22"),
    ]
    for name, page in items:
        p(doc, f"{name}\t{page}", align=WD_ALIGN_PARAGRAPH.LEFT, indent=False)
    doc.add_page_break()


def add_intro(doc):
    heading(doc, "Введение", 1)
    p(
        doc,
        "Целью курсовой работы является разработка учебной системы управления базами данных на языке C++ "
        "с поддержкой целочисленных и строковых данных, SQL-подобного языка запросов, файлового хранения "
        "и индексирования с помощью B+-дерева. Разработка выполнялась как консольное приложение, которое "
        "может работать в интерактивном и пакетном режимах.",
    )
    p(
        doc,
        "В рамках работы реализованы уровни системы, базы данных и таблицы. Каждая база данных содержит "
        "набор таблиц, а каждая таблица хранит схему, записи и индексы для столбцов с модификатором INDEXED. "
        "Для разбора языка запросов используются генераторы Flex и Bison, для формирования JSON-результатов "
        "используется библиотека nlohmann/json.",
    )
    p(
        doc,
        "Дополнительно реализованы шесть расширений: темпоральная персистентность для отката таблицы, "
        "интернирование строк, журналирование запросов, значения по умолчанию, составные логические "
        "выражения в WHERE и агрегатные функции SELECT.",
    )


def add_base_section(doc):
    heading(doc, "Основная часть", 1)
    heading(doc, "1. Базовая СУБД с индексом B+-tree (задание 0)", 2)
    p(
        doc,
        "Базовая часть проекта включает иерархию хранения, синтаксический анализ SQL-подобного языка, "
        "семантическую проверку запросов, выполнение DDL- и DML-операций, сохранение данных в файловой "
        "системе и использование B+-tree индекса для ускорения поиска.",
    )

    heading(doc, "1.1. Архитектура программы", 3)
    p(
        doc,
        "Точка входа программы находится в файле src/main.cpp и содержит только вызов функции "
        "course_dbms_main. Основная логика вынесена в модуль src/dbms.cpp, что отделяет запуск приложения "
        "от работы с каталогом баз данных, таблицами, индексами и командами. Общие структуры абстрактного "
        "синтаксического дерева расположены в include/db_ast.h, а публичный интерфейс парсера — в "
        "include/sql_parser_api.h.",
    )
    p(
        doc,
        "Такое разделение упрощает защиту проекта: main отвечает за запуск, Flex/Bison отвечают за перевод "
        "текста запроса в структуру Command, а dbms.cpp отвечает за проверку и выполнение команд.",
    )

    heading(doc, "1.2. Лексический и синтаксический анализ", 3)
    p(
        doc,
        "Лексер, описанный в parser/sql_lexer.l, разбивает входной текст на токены: ключевые слова, "
        "идентификаторы, числа, строковые литералы и знаки операций. Парсер, описанный в parser/sql_parser.y, "
        "проверяет порядок токенов и строит объект команды. Если синтаксис нарушен, пользователь получает "
        "сообщение об ошибке, а программа продолжает работу.",
    )
    p(
        doc,
        "Использование Flex/Bison уменьшает количество ручного кода для распознавания запросов и делает "
        "грамматику языка отдельной частью проекта. Это соответствует рекомендации преподавателя по "
        "использованию готовых средств лексического и синтаксического анализа.",
    )

    heading(doc, "1.3. Хранение данных", 3)
    p(
        doc,
        "Все данные сохраняются в каталоге db_data. Для каждой базы данных создаётся отдельная директория, "
        "для каждой таблицы — отдельная директория внутри базы. Схема таблицы хранится в schema.json, "
        "строки таблицы — в rows.dat. При запуске запросов таблица загружается из файлов, после изменения "
        "данные записываются обратно.",
    )
    p(
        doc,
        "В таблице поддерживаются типы int и string, а также значение NULL. Модификатор NOT_NULL запрещает "
        "отсутствие значения, а INDEXED одновременно делает столбец обязательным, уникальным и индексируемым. "
        "При вставке и обновлении выполняется проверка типа, уникальности и ограничений целостности.",
    )

    heading(doc, "1.4. Индекс на основе B+-дерева", 3)
    p(
        doc,
        "Для индексируемых столбцов автоматически создаётся B+-tree индекс. В индексе ключом является "
        "значение столбца, а значением — идентификатор строки. Сами записи таблицы в индекс не копируются, "
        "поэтому выполняется требование о запрете дублирования данных при построении индексов.",
    )
    p(
        doc,
        "При выполнении SELECT, UPDATE и DELETE система пытается извлечь из условия простой индексируемый "
        "предикат. Для операций равенства и диапазонных сравнений по индексируемому столбцу используется "
        "поиск по B+-tree, после чего оставшаяся часть условия дополнительно проверяется на найденных строках. "
        "Если подходящего индекса нет, выполняется полный просмотр таблицы.",
    )

    heading(doc, "1.5. Поддерживаемые команды", 3)
    p(
        doc,
        "Реализованы команды CREATE DATABASE, DROP DATABASE, USE, CREATE TABLE, DROP TABLE, INSERT, UPDATE, "
        "DELETE и SELECT. Команды могут вводиться многострочно и завершаются символом точки с запятой. "
        "SELECT возвращает массив объектов JSON; для остальных успешных команд выводится статус OK.",
    )
    p(
        doc,
        "Таблицы могут указываться через активный контекст USE либо в формате database.table. В SELECT "
        "поддерживается выбор всех столбцов или заданного списка столбцов с алиасами через AS.",
    )


def add_extra(doc, number, title, algorithm, implementation, result):
    heading(doc, f"{number}. {title}", 2)
    p(doc, "Алгоритм.", bold=True)
    p(doc, algorithm)
    p(doc, "Реализация.", bold=True)
    p(doc, implementation)
    p(doc, "Результат.", bold=True)
    p(doc, result)


def add_extras(doc):
    add_extra(
        doc,
        "2",
        "Поддержка темпоральной персистентности (задание 1)",
        "Перед изменением таблицы система сохраняет в журнале обратимую информацию об операции. Для INSERT "
        "достаточно знать добавленные строки, для DELETE — удалённые строки, для UPDATE — состояние строк до "
        "обновления. При REVERT журнал просматривается в обратном порядке, а изменения после указанного "
        "момента отменяются.",
        "Для каждой таблицы используется файл history.log. Команда REVERT table yyyy.mm.dd-hh:mm:ss.msmsms "
        "не копирует весь каталог базы данных, а применяет обратные операции к данным конкретной таблицы. "
        "После отката индексы перестраиваются по актуальным строкам.",
        "Реализован откат состояния таблицы к указанному моменту времени без snapshot-копирования файлов.",
    )
    add_extra(
        doc,
        "3",
        "Оптимизация хранения строковых данных (задание 2)",
        "При появлении строкового значения система проверяет пул уже известных строк. Если такая строка уже "
        "существует, новое значение заменяется ссылкой на существующий объект. Если строки ещё нет, она "
        "добавляется в пул.",
        "В include/db_ast.h тип строкового значения представлен как shared_ptr<const string>. Функция "
        "intern_string возвращает общий указатель на единственный экземпляр строки. Сравнение значений при "
        "этом остаётся лексикографическим по содержимому строки.",
        "Одинаковые строковые значения в оперативной памяти хранятся в единственном экземпляре.",
    )
    add_extra(
        doc,
        "4",
        "Логирование активности (задание 7)",
        "При обработке каждого запроса фиксируются тело запроса, идентификатор клиента, идентификатор "
        "обработчика, время начала, время завершения, длительность обработки и итоговый статус.",
        "Журнал записывается в db_data/access.log. В консольном режиме идентификатор клиента имеет значение "
        "local-cli, а идентификатор обработчика соответствует текущему потоку выполнения.",
        "Получен файл доступа, по которому можно восстановить историю запросов и увидеть ошибки выполнения.",
    )
    add_extra(
        doc,
        "5",
        "Значения по умолчанию (задание 10)",
        "При создании таблицы столбец может получить модификатор DEFAULT. Если при INSERT значение для "
        "такого столбца не передано, используется заданное значение по умолчанию. Если DEFAULT нет, а "
        "столбец допускает NULL, записывается NULL.",
        "Грамматика CREATE TABLE расширена правилом DEFAULT value. Значение сохраняется в схеме таблицы и "
        "учитывается при сборке полной строки перед вставкой.",
        "INSERT стал удобнее: пользователь может передавать только те столбцы, которые отличаются от "
        "значений по умолчанию.",
    )
    add_extra(
        doc,
        "6",
        "Составные логические выражения (задание 11)",
        "Условие WHERE представлено деревом выражений. Листья дерева — сравнения, BETWEEN и LIKE, внутренние "
        "узлы — логические операции AND и OR. Для скобок создаётся вложенное поддерево, что сохраняет "
        "приоритет вычисления.",
        "В Bison-грамматике задан приоритет: сначала обрабатываются выражения в скобках, затем AND, затем OR. "
        "Во время выполнения условие рекурсивно вычисляется для строки таблицы.",
        "Появилась поддержка запросов вида WHERE (age >= 18 AND city == \"Moscow\") OR name LIKE \"A.*\".",
    )
    add_extra(
        doc,
        "7",
        "Агрегатные функции SUM, COUNT, AVG (задание 12)",
        "После отбора строк по WHERE система проходит по выбранным значениям и вычисляет агрегат. COUNT "
        "считает количество подходящих строк, SUM складывает числовые значения, AVG делит сумму на количество "
        "не NULL-значений.",
        "Парсер распознаёт SUM(column), COUNT(column) и AVG(column) внутри SELECT. Результат возвращается "
        "как JSON-объект с именем агрегата или заданным алиасом.",
        "Реализованы базовые аналитические запросы без необходимости обрабатывать результат SELECT вручную.",
    )


def add_testing(doc):
    heading(doc, "Тестирование", 1)
    p(
        doc,
        "Проверка выполнялась сборкой проекта через CMake и запуском пакетного сценария. Тестовый сценарий "
        "создаёт базу данных, таблицу с индексируемым столбцом, вставляет строки, выполняет SELECT, UPDATE, "
        "DELETE, проверяет ошибку уникальности INDEXED, значения DEFAULT, агрегаты и откат REVERT.",
    )
    listing_caption(doc, 1, "Пример сценария для проверки базовых операций")
    code(
        doc,
        r'''
CREATE DATABASE demo;
USE demo;
CREATE TABLE products (
    id int INDEXED,
    title string NOT_NULL,
    qty int DEFAULT 0
);
INSERT INTO products (id, title, qty) VALUE (1, "one", 10), (2, "two", 20);
SELECT * FROM products WHERE id BETWEEN 1 AND 3;
UPDATE products SET qty = 99 WHERE id == 1;
SELECT COUNT(id) AS total, SUM(qty) AS qty_sum FROM products;
''',
    )
    p(
        doc,
        "После выполнения SELECT программа выводит массив JSON-объектов. Ошибочные запросы не приводят к "
        "аварийному завершению: система возвращает диагностическое сообщение, например о неизвестной таблице, "
        "нарушении NOT_NULL или повторном значении в индексируемом столбце.",
    )
    p(
        doc,
        "Также проверялась сборка проекта командой cmake --build build -j2. После рефакторинга main.cpp "
        "остался минимальным, а основной модуль dbms.cpp успешно компилируется вместе с файлами, "
        "сгенерированными Flex и Bison.",
    )


def add_conclusion_sources_appendix(doc):
    heading(doc, "Вывод", 1)
    p(
        doc,
        "В результате выполнения курсовой работы разработана консольная СУБД на C++ с файловым хранением, "
        "типизированными таблицами, ограничениями целостности и индексом на основе B+-дерева. Реализован "
        "SQL-подобный язык запросов, который разбирается средствами Flex и Bison. Для выборок результат "
        "формируется в формате JSON.",
    )
    p(
        doc,
        "Помимо обязательной части реализованы шесть дополнительных заданий: темпоральная персистентность, "
        "интернирование строк, логирование активности, DEFAULT-значения, составные WHERE-условия и агрегатные "
        "функции. Проект имеет модульную структуру, что упрощает дальнейшее расширение, например добавление "
        "клиент-серверного режима или более строгого механизма восстановления после сбоев.",
    )

    heading(doc, "Список использованных источников", 1)
    sources = [
        "1. ISO/IEC 14882:2020. Programming languages — C++.",
        "2. Aho A. V., Lam M. S., Sethi R., Ullman J. D. Compilers: Principles, Techniques, and Tools. 2nd ed. Addison-Wesley, 2006.",
        "3. Bayer R., McCreight E. Organization and Maintenance of Large Ordered Indexes. Acta Informatica, 1972.",
        "4. Flex: The Fast Lexical Analyzer. URL: https://github.com/westes/flex",
        "5. GNU Bison Manual. URL: https://www.gnu.org/software/bison/manual/",
        "6. nlohmann/json: JSON for Modern C++. URL: https://github.com/nlohmann/json",
    ]
    for source in sources:
        p(doc, source, align=WD_ALIGN_PARAGRAPH.LEFT, indent=False)

    doc.add_page_break()
    heading(doc, "Приложение А. Пример сценария работы", 1)
    listing_caption(doc, 2, "Расширенный демонстрационный сценарий")
    code(
        doc,
        r'''
CREATE DATABASE course;
USE course;
CREATE TABLE students (
    id int INDEXED,
    name string NOT_NULL,
    group_name string DEFAULT "M8O",
    age int DEFAULT 18
);
INSERT INTO students (id, name, age) VALUE (1, "Alex", 19), (2, "Maria", 20);
SELECT name AS student_name, age FROM students
WHERE (age >= 18 AND group_name == "M8O") OR name LIKE "A.*";
SELECT COUNT(id) AS total, AVG(age) AS avg_age FROM students;
REVERT students 2026.05.17-20:00:00.000000;
''',
    )

    heading(doc, "Приложение Б. Ключевые фрагменты реализации", 1)
    listing_caption(doc, 3, "Минимальная точка входа приложения")
    code(
        doc,
        r'''
#include <app.h>

int main(int argc, char** argv) {
    return course_dbms_main(argc, argv);
}
''',
    )
    listing_caption(doc, 4, "Идея хранения строк через интернирование")
    code(
        doc,
        r'''
using InternedString = std::shared_ptr<const std::string>;

inline InternedString intern_string(const std::string& value) {
    static std::unordered_map<std::string, std::weak_ptr<const std::string>> pool;
    auto found = pool.find(value);
    if (found != pool.end()) {
        if (auto existing = found->second.lock()) {
            return existing;
        }
    }
    auto created = std::make_shared<const std::string>(value);
    pool[value] = created;
    return created;
}
''',
    )


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = configure_document()
    add_title_page(doc)
    add_contents(doc)
    add_intro(doc)
    add_base_section(doc)
    add_extras(doc)
    add_testing(doc)
    add_conclusion_sources_appendix(doc)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
